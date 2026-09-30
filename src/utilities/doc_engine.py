"""
EFFIONG AI - Sovereign Knowledge Document Channel (native high-fidelity asset engine)
=====================================================================================
Compiles Markdown produced by the brain into REAL documents:

  compile_pdf_bytes   - PDF with a strict typographic hierarchy, cover band, running header, "Page x of y",
                        wrapped tables with a colour-coded schema (status words shade green / amber / red)
  compile_word_bytes  - Word (.docx) with real heading styles, shaded tables, lists, header/footer page numbers
  compile_pptx_bytes  - PowerPoint deck in the Effiong dark brand theme (one slide per "## " heading)
  compile_certificate_pdf - landscape certificate with borders, seal and signature lines
  render_chart        - PNG charts (bar, line, area, pie, scatter) from a JSON spec
  sanitize_svg        - makes vector graphics (SVG, importable into CorelDRAW / Illustrator / Inkscape) safe to show

Markdown supported: # headings, paragraphs, **bold**, *italic*, `code`, links, bullet / numbered lists (nested),
pipe tables, block quotes, code fences, horizontal rules and ![alt](image-name) for supplied images.

Every builder is defensive: a malformed table or odd character can never stop a document from being produced.
"""
from __future__ import annotations

import io
import os
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from xml.sax.saxutils import escape as xml_escape

from src.core.health import HEALTH, safe_import
from src.utilities.text_utils import strip_emoji

BRAND_ORANGE = "#D27D2D"
BRAND_DARK = "#0B0F17"
BRAND_PANEL = "#121620"
TEXT_DARK = "#1F2937"
TEXT_MUTED = "#6B7280"
ZEBRA = "#F3F4F6"
GRID = "#D1D5DB"

GREEN_WORDS = {"verified", "approved", "pass", "passed", "complete", "completed", "done", "active", "yes", "low", "safe",
               "confirmed", "valid", "success", "achieved", "on track", "fact", "evidence supported"}
AMBER_WORDS = {"pending", "review", "in review", "partial", "medium", "moderate", "oral", "probable", "warning", "draft",
               "in progress", "pending verification", "oral tradition", "watch", "unknown", "speculation"}
RED_WORDS = {"failed", "fail", "rejected", "high", "critical", "risk", "no", "expired", "disputed", "unverified", "overdue",
             "blocked", "denied", "at risk", "false", "extreme"}


def _status_color(cell_text: str) -> Optional[Tuple[str, str]]:
    """(background, text) colours for a status-like cell, else None."""
    t = re.sub(r"[*_`]", "", cell_text or "").strip().lower()
    if t in GREEN_WORDS:
        return "#DCFCE7", "#166534"
    if t in AMBER_WORDS:
        return "#FEF3C7", "#92400E"
    if t in RED_WORDS:
        return "#FEE2E2", "#991B1B"
    return None


# =========================================================================================
# Markdown parsing
# =========================================================================================
Block = Dict[str, Any]
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _split_row(line: str) -> List[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def parse_markdown(text: str) -> List[Block]:
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    blocks: List[Block] = []
    para: List[str] = []
    i = 0

    def flush_para() -> None:
        if para:
            blocks.append({"type": "para", "text": " ".join(s.strip() for s in para).strip()})
            para.clear()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_para()
            i += 1
            code: List[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append({"type": "code", "text": "\n".join(code)})
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            flush_para()
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip().rstrip("#").strip()})
            i += 1
            continue

        if re.match(r"^([-*_])\1{2,}$", stripped.replace(" ", "")):
            flush_para()
            blocks.append({"type": "rule"})
            i += 1
            continue

        if stripped == "\\pagebreak" or stripped == "<pagebreak>":
            flush_para()
            blocks.append({"type": "pagebreak"})
            i += 1
            continue

        # table: a row followed by a separator row
        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            flush_para()
            header = _split_row(stripped)
            i += 2
            rows: List[List[str]] = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i]))
                i += 1
            width = len(header)
            rows = [(r + [""] * width)[:width] for r in rows]
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        m_img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if m_img:
            flush_para()
            blocks.append({"type": "image", "alt": m_img.group(1), "name": m_img.group(2)})
            i += 1
            continue

        m_b = re.match(r"^(\s*)([-*•+])\s+(.*)$", line)
        m_n = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if m_b or m_n:
            flush_para()
            ordered = bool(m_n)
            items: List[Tuple[int, str]] = []
            while i < len(lines):
                lb = re.match(r"^(\s*)([-*•+])\s+(.*)$", lines[i])
                ln = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", lines[i])
                mm = ln if ordered else lb
                if not mm:
                    break
                indent = len(mm.group(1).replace("\t", "    "))
                items.append((min(2, indent // 2), mm.group(3).strip()))
                i += 1
                # soft-wrapped continuation lines
                while i < len(lines) and lines[i].strip() and not re.match(r"^\s*([-*•+]|\d+[.)])\s+", lines[i]) \
                        and not lines[i].strip().startswith(("#", ">", "|", "```")):
                    items[-1] = (items[-1][0], items[-1][1] + " " + lines[i].strip())
                    i += 1
            blocks.append({"type": "numbered" if ordered else "bullets", "items": items})
            continue

        if stripped.startswith(">"):
            flush_para()
            quote: List[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append({"type": "quote", "text": " ".join(quote)})
            continue

        para.append(stripped)
        i += 1
    flush_para()
    return blocks


_INLINE = re.compile(r"(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|__(.+?)__|\*(.+?)\*|_(.+?)_|`(.+?)`|\[([^\]]+)\]\(([^)]+)\))")


def parse_inline(text: str) -> List[Tuple[str, Dict[str, Any]]]:
    """Split inline markdown into (text, flags) runs. flags: b, i, code, link."""
    runs: List[Tuple[str, Dict[str, Any]]] = []
    pos = 0
    text = text or ""
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], {}))
        if m.group(2):
            runs.append((m.group(2), {"b": True, "i": True}))
        elif m.group(3) or m.group(4):
            runs.append((m.group(3) or m.group(4), {"b": True}))
        elif m.group(5) or m.group(6):
            runs.append((m.group(5) or m.group(6), {"i": True}))
        elif m.group(7):
            runs.append((m.group(7), {"code": True}))
        else:
            runs.append((m.group(8), {"link": m.group(9)}))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], {}))
    return runs or [("", {})]


def plain(text: str) -> str:
    return "".join(t for t, _ in parse_inline(text))


def extract_title(blocks: List[Block], fallback: str = "Document") -> Tuple[str, List[Block]]:
    """If the first block is a level-1 heading, use it as the document title and drop it from the body."""
    for idx, b in enumerate(blocks):
        if b["type"] == "heading" and b["level"] == 1:
            return plain(b["text"]), blocks[:idx] + blocks[idx + 1:]
        if b["type"] != "rule":
            break
    return fallback, blocks


def _clean(text: str) -> str:
    """Remove characters the built-in document fonts cannot draw (emoji etc.)."""
    return strip_emoji(text or "").replace("\u200b", "")


# =========================================================================================
# PDF
# =========================================================================================
_FONTS_READY: Optional[bool] = None
F_REG, F_BOLD, F_ITAL, F_BI, F_MONO = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique", "Courier"


def _register_fonts() -> None:
    """Use DejaVu (ships with matplotlib) so African diacritics / Greek / Cyrillic render; fall back to Helvetica."""
    global _FONTS_READY, F_REG, F_BOLD, F_ITAL, F_BI, F_MONO
    if _FONTS_READY is not None:
        return
    _FONTS_READY = False
    try:
        import matplotlib
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        base = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
        files = {"EffSans": "DejaVuSans.ttf", "EffSans-Bold": "DejaVuSans-Bold.ttf",
                 "EffSans-Italic": "DejaVuSans-Oblique.ttf", "EffSans-BoldItalic": "DejaVuSans-BoldOblique.ttf",
                 "EffMono": "DejaVuSansMono.ttf"}
        for name, fname in files.items():
            pdfmetrics.registerFont(TTFont(name, os.path.join(base, fname)))
        pdfmetrics.registerFontFamily("EffSans", normal="EffSans", bold="EffSans-Bold", italic="EffSans-Italic",
                                      boldItalic="EffSans-BoldItalic")
        F_REG, F_BOLD, F_ITAL, F_BI, F_MONO = "EffSans", "EffSans-Bold", "EffSans-Italic", "EffSans-BoldItalic", "EffMono"
        _FONTS_READY = True
    except Exception as exc:
        HEALTH.flag("pdf-fonts", f"DejaVu unavailable, using Helvetica: {exc}")


def _rl_markup(text: str) -> str:
    """Inline markdown -> reportlab paragraph markup (escaped)."""
    out = []
    for run, fl in parse_inline(_clean(text)):
        piece = xml_escape(run)
        if fl.get("code"):
            piece = f'<font name="{F_MONO}" size="9" color="#B45309">{piece}</font>'
        if fl.get("b"):
            piece = f"<b>{piece}</b>"
        if fl.get("i"):
            piece = f"<i>{piece}</i>"
        if fl.get("link"):
            piece = f'<a href="{xml_escape(fl["link"], {chr(34): "&quot;"})}" color="#2563EB">{piece}</a>'
        out.append(piece)
    return "".join(out)


def compile_pdf_bytes(markdown_text: str, title: Optional[str] = None, subtitle: str = "", doc_type: str = "Document",
                      images: Optional[Dict[str, bytes]] = None, author: str = "Effiong AI") -> bytes:
    """Markdown -> professional PDF (bytes). Never raises for content problems."""
    _register_fonts()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (HRFlowable, Image as RLImage, KeepTogether, PageBreak, Paragraph, Preformatted,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    blocks = parse_markdown(markdown_text)
    doc_title, blocks = extract_title(blocks, title or "Effiong AI Document")
    doc_title = _clean(title or doc_title)
    images = images or {}
    now_str = datetime.now().strftime("%d %B %Y, %H:%M")

    S = {
        "title": ParagraphStyle("t", fontName=F_BOLD, fontSize=25, leading=30, textColor=colors.HexColor(BRAND_DARK), spaceAfter=6),
        "subtitle": ParagraphStyle("st", fontName=F_REG, fontSize=12, leading=16, textColor=colors.HexColor(TEXT_MUTED), spaceAfter=14),
        "h1": ParagraphStyle("h1", fontName=F_BOLD, fontSize=17, leading=22, textColor=colors.HexColor(BRAND_DARK), spaceBefore=16, spaceAfter=4),
        "h2": ParagraphStyle("h2", fontName=F_BOLD, fontSize=13.5, leading=18, textColor=colors.HexColor(BRAND_ORANGE), spaceBefore=12, spaceAfter=3),
        "h3": ParagraphStyle("h3", fontName=F_BOLD, fontSize=11.5, leading=15, textColor=colors.HexColor(TEXT_DARK), spaceBefore=8, spaceAfter=2),
        "h4": ParagraphStyle("h4", fontName=F_BI, fontSize=10.5, leading=14, textColor=colors.HexColor(TEXT_DARK), spaceBefore=6, spaceAfter=2),
        "body": ParagraphStyle("b", fontName=F_REG, fontSize=10.5, leading=15.5, textColor=colors.HexColor(TEXT_DARK), spaceAfter=6, alignment=TA_LEFT),
        "quote": ParagraphStyle("q", fontName=F_ITAL, fontSize=10.5, leading=15, textColor=colors.HexColor("#4B5563"), leftIndent=14,
                                borderPadding=(6, 6, 6, 8), backColor=colors.HexColor("#FFF7ED"), spaceBefore=6, spaceAfter=8),
        "cell": ParagraphStyle("c", fontName=F_REG, fontSize=9, leading=12, textColor=colors.HexColor(TEXT_DARK)),
        "cellh": ParagraphStyle("ch", fontName=F_BOLD, fontSize=9, leading=12, textColor=colors.white),
        "caption": ParagraphStyle("cap", fontName=F_ITAL, fontSize=8.5, leading=11, textColor=colors.HexColor(TEXT_MUTED), alignment=TA_CENTER, spaceAfter=8),
    }

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *a: Any, **k: Any) -> None:
            super().__init__(*a, **k)
            self._saved: List[dict] = []

        def showPage(self) -> None:
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self) -> None:
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._draw_chrome(total)
                super().showPage()
            super().save()

        def _draw_chrome(self, total: int) -> None:
            w, h = A4
            page = self._pageNumber
            if page == 1:
                self.setFillColor(colors.HexColor(BRAND_DARK))
                self.rect(0, h - 26 * mm, w, 26 * mm, stroke=0, fill=1)
                self.setFillColor(colors.HexColor(BRAND_ORANGE))
                self.setFont(F_BOLD, 13)
                self.drawString(20 * mm, h - 13 * mm, "EFFIONG AI")
                self.setFillColor(colors.HexColor("#8B949E"))
                self.setFont(F_REG, 8)
                self.drawString(20 * mm, h - 19 * mm, "SOVEREIGN WISDOM ENGINE")
                self.setFillColor(colors.HexColor("#F0F6FC"))
                self.setFont(F_BOLD, 9)
                self.drawRightString(w - 20 * mm, h - 13 * mm, _clean(doc_type).upper()[:40])
                self.setFont(F_REG, 8)
                self.setFillColor(colors.HexColor("#8B949E"))
                self.drawRightString(w - 20 * mm, h - 19 * mm, now_str)
                self.setFillColor(colors.HexColor(BRAND_ORANGE))
                self.rect(0, h - 27.2 * mm, w, 1.2 * mm, stroke=0, fill=1)
            else:
                self.setStrokeColor(colors.HexColor(BRAND_ORANGE))
                self.setLineWidth(1.2)
                self.line(20 * mm, h - 14 * mm, w - 20 * mm, h - 14 * mm)
                self.setFillColor(colors.HexColor(TEXT_MUTED))
                self.setFont(F_REG, 8)
                self.drawString(20 * mm, h - 11 * mm, doc_title[:90])
                self.drawRightString(w - 20 * mm, h - 11 * mm, "EFFIONG AI")
            self.setStrokeColor(colors.HexColor(GRID))
            self.setLineWidth(0.5)
            self.line(20 * mm, 14 * mm, w - 20 * mm, 14 * mm)
            self.setFillColor(colors.HexColor(TEXT_MUTED))
            self.setFont(F_REG, 8)
            self.drawString(20 * mm, 9.5 * mm, "Generated by Effiong AI - verify critical facts before relying on this document")
            self.drawRightString(w - 20 * mm, 9.5 * mm, f"Page {page} of {total}")

    story: List[Any] = [Spacer(1, 6 * mm), Paragraph(xml_escape(doc_title), S["title"])]
    if subtitle:
        story.append(Paragraph(_rl_markup(subtitle), S["subtitle"]))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor(BRAND_ORANGE), spaceAfter=8))

    avail_w = A4[0] - 40 * mm

    def make_table(block: Block) -> Any:
        header, rows = block["header"], block["rows"]
        ncol = max(1, len(header))
        lens = [max([len(plain(header[c]))] + [len(plain(r[c])) for r in rows] + [4]) for c in range(ncol)]
        weights = [min(max(l, 8), 40) for l in lens]
        total = float(sum(weights))
        widths = [avail_w * w / total for w in weights]
        data = [[Paragraph(_rl_markup(h), S["cellh"]) for h in header]]
        style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_DARK)),
            ("LINEBELOW", (0, 0), (-1, 0), 1.6, colors.HexColor(BRAND_ORANGE)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(GRID)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]
        for r_i, row in enumerate(rows, start=1):
            cells = []
            if r_i % 2 == 0:
                style_cmds.append(("BACKGROUND", (0, r_i), (-1, r_i), colors.HexColor(ZEBRA)))
            for c_i, cell in enumerate(row):
                col = _status_color(cell)
                if col:
                    style_cmds.append(("BACKGROUND", (c_i, r_i), (c_i, r_i), colors.HexColor(col[0])))
                    cells.append(Paragraph(f'<font color="{col[1]}"><b>{_rl_markup(cell)}</b></font>', S["cell"]))
                else:
                    cells.append(Paragraph(_rl_markup(cell), S["cell"]))
            data.append(cells)
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle(style_cmds))
        return t

    for b in blocks:
        try:
            t = b["type"]
            if t == "heading":
                lvl = min(4, max(1, b["level"]))
                # '#'/'##' are major sections, '###' sub-sections, '####' minor headings
                key = {1: "h1", 2: "h1", 3: "h2", 4: "h3"}[lvl]
                para = Paragraph(_rl_markup(b["text"]), S[key])
                if lvl <= 2:
                    story.append(KeepTogether([para, HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(GRID), spaceAfter=4)]))
                else:
                    story.append(para)
            elif t == "para":
                story.append(Paragraph(_rl_markup(b["text"]), S["body"]))
            elif t in ("bullets", "numbered"):
                for n, (level, text) in enumerate(b["items"], 1):
                    st = ParagraphStyle("li", parent=S["body"], leftIndent=16 + level * 14, bulletIndent=4 + level * 14, spaceAfter=2.5)
                    bullet = f"{n}." if t == "numbered" else ("•" if level == 0 else "–")
                    story.append(Paragraph(_rl_markup(text), st, bulletText=bullet))
                story.append(Spacer(1, 4))
            elif t == "table":
                story.append(make_table(b))
                story.append(Spacer(1, 8))
            elif t == "quote":
                story.append(Paragraph(_rl_markup(b["text"]), S["quote"]))
            elif t == "code":
                story.append(Preformatted(_clean(b["text"])[:4000], ParagraphStyle("code", fontName=F_MONO, fontSize=8.5, leading=11,
                                                                                  backColor=colors.HexColor("#F3F4F6"), borderPadding=6,
                                                                                  leftIndent=6, spaceAfter=8)))
            elif t == "rule":
                story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(GRID), spaceBefore=6, spaceAfter=6))
            elif t == "pagebreak":
                story.append(PageBreak())
            elif t == "image" and b["name"] in images:
                img = RLImage(io.BytesIO(images[b["name"]]))
                iw, ih = img.imageWidth, img.imageHeight
                scale = min(avail_w / iw, (110 * mm) / ih, 1.0 if iw < avail_w else avail_w / iw)
                img.drawWidth, img.drawHeight = iw * scale, ih * scale
                story.append(KeepTogether([img, Paragraph(_rl_markup(b["alt"]), S["caption"])]))
        except Exception as exc:  # one bad block must not lose the whole document
            HEALTH.flag("pdf-build", f"block {b.get('type')}: {exc.__class__.__name__}: {exc}")
            story.append(Paragraph(xml_escape(plain(str(b.get("text", "")))[:500]), S["body"]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=32 * mm, bottomMargin=20 * mm,
                            title=doc_title, author=author, subject=doc_type)
    doc.build(story, canvasmaker=NumberedCanvas)
    return buf.getvalue()


# =========================================================================================
# Certificate (landscape)
# =========================================================================================
def compile_certificate_pdf(title: str, recipient: str, body: str, issuer: str = "Effiong AI", date_text: str = "",
                            serial: str = "", signatories: Optional[List[Tuple[str, str]]] = None) -> bytes:
    _register_fonts()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as rl_canvas

    buf = io.BytesIO()
    w, h = landscape(A4)
    c = rl_canvas.Canvas(buf, pagesize=landscape(A4))
    c.setTitle(_clean(title))
    c.setFillColor(colors.HexColor("#FFFDF8"))
    c.rect(0, 0, w, h, stroke=0, fill=1)
    c.setStrokeColor(colors.HexColor(BRAND_DARK))
    c.setLineWidth(3)
    c.rect(10 * mm, 10 * mm, w - 20 * mm, h - 20 * mm)
    c.setStrokeColor(colors.HexColor(BRAND_ORANGE))
    c.setLineWidth(1.4)
    c.rect(14 * mm, 14 * mm, w - 28 * mm, h - 28 * mm)

    c.setFillColor(colors.HexColor(BRAND_ORANGE))
    c.setFont(F_BOLD, 12)
    c.drawCentredString(w / 2, h - 30 * mm, "EFFIONG AI  -  SOVEREIGN WISDOM ENGINE")
    c.setFillColor(colors.HexColor(BRAND_DARK))
    c.setFont(F_BOLD, 34)
    c.drawCentredString(w / 2, h - 52 * mm, _clean(title).upper()[:44])
    c.setFillColor(colors.HexColor(TEXT_MUTED))
    c.setFont(F_ITAL, 13)
    c.drawCentredString(w / 2, h - 66 * mm, "This is to certify that")
    c.setFillColor(colors.HexColor(BRAND_ORANGE))
    c.setFont(F_BOLD, 30)
    c.drawCentredString(w / 2, h - 84 * mm, _clean(recipient)[:50])
    c.setStrokeColor(colors.HexColor(BRAND_DARK))
    c.setLineWidth(0.8)
    c.line(w / 2 - 70 * mm, h - 88 * mm, w / 2 + 70 * mm, h - 88 * mm)

    from reportlab.lib.utils import simpleSplit

    lines = simpleSplit(_clean(body), F_REG, 12, w - 90 * mm)[:7]
    c.setFillColor(colors.HexColor(TEXT_DARK))
    c.setFont(F_REG, 12)
    y = h - 102 * mm
    for ln in lines:
        c.drawCentredString(w / 2, y, ln)
        y -= 7 * mm

    # seal
    cx, cy = w / 2, 42 * mm
    c.setFillColor(colors.HexColor(BRAND_ORANGE))
    c.circle(cx, cy, 15 * mm, stroke=0, fill=1)
    c.setStrokeColor(colors.white)
    c.setLineWidth(1)
    c.circle(cx, cy, 12.5 * mm, stroke=1, fill=0)
    c.setFillColor(colors.white)
    c.setFont(F_BOLD, 8)
    c.drawCentredString(cx, cy + 2 * mm, "EFFIONG")
    c.drawCentredString(cx, cy - 3 * mm, "AI SEAL")

    sigs = signatories or [(issuer, "Issuer"), ("Authorised Signatory", "Verification")]
    for idx, (name, role) in enumerate(sigs[:2]):
        x = 55 * mm if idx == 0 else w - 55 * mm
        c.setStrokeColor(colors.HexColor(TEXT_DARK))
        c.setLineWidth(0.7)
        c.line(x - 34 * mm, 33 * mm, x + 34 * mm, 33 * mm)
        c.setFillColor(colors.HexColor(TEXT_DARK))
        c.setFont(F_BOLD, 10)
        c.drawCentredString(x, 27.5 * mm, _clean(name)[:36])
        c.setFont(F_REG, 8.5)
        c.setFillColor(colors.HexColor(TEXT_MUTED))
        c.drawCentredString(x, 22.5 * mm, _clean(role)[:36])
    c.setFont(F_REG, 8.5)
    c.setFillColor(colors.HexColor(TEXT_MUTED))
    c.drawString(20 * mm, 18 * mm, f"Date: {date_text or datetime.now().strftime('%d %B %Y')}")
    if serial:
        c.drawRightString(w - 20 * mm, 18 * mm, f"Certificate No: {serial}")
    c.showPage()
    c.save()
    return buf.getvalue()


# =========================================================================================
# DOCX
# =========================================================================================
def _shade(cell: Any, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    tcPr.append(shd)


def _add_runs(paragraph: Any, text: str, size: Optional[float] = None, color: Optional[str] = None, bold: Optional[bool] = None) -> None:
    from docx.shared import Pt, RGBColor

    for run_text, fl in parse_inline(_clean(text)):
        run = paragraph.add_run(run_text)
        run.bold = True if (fl.get("b") or bold) else None
        run.italic = True if fl.get("i") else None
        if fl.get("code"):
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
        else:
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        if fl.get("link"):
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            run.font.underline = True


def _docx_page_number(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for code in ("PAGE",):
        run = paragraph.add_run()
        for kind, txt in (("begin", None), (None, code), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = txt
            run._r.append(el)


def compile_word_bytes(markdown_text: str, title: Optional[str] = None, subtitle: str = "", doc_type: str = "Document",
                       images: Optional[Dict[str, bytes]] = None, author: str = "Effiong AI") -> bytes:
    """Markdown -> Word document (bytes)."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Pt, RGBColor

    blocks = parse_markdown(markdown_text)
    doc_title, blocks = extract_title(blocks, title or "Effiong AI Document")
    doc_title = _clean(title or doc_title)
    images = images or {}

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)
    for name, size, color in (("Heading 1", 18, BRAND_DARK), ("Heading 2", 14, BRAND_ORANGE), ("Heading 3", 12, TEXT_DARK), ("Heading 4", 11, TEXT_DARK)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    doc.core_properties.title = doc_title
    doc.core_properties.author = author
    doc.core_properties.subject = doc_type

    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("EFFIONG AI  |  " + doc_title[:70])
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(BRAND_ORANGE.lstrip("#"))
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Generated by Effiong AI  -  Page ")
    fr.font.size = Pt(8.5)
    _docx_page_number(fp)

    tp = doc.add_paragraph()
    tr = tp.add_run(doc_title)
    tr.bold = True
    tr.font.size = Pt(26)
    tr.font.color.rgb = RGBColor.from_string(BRAND_DARK.lstrip("#"))
    meta = doc.add_paragraph()
    mr = meta.add_run(f"{_clean(doc_type)}  -  {datetime.now().strftime('%d %B %Y')}")
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor.from_string(TEXT_MUTED.lstrip("#"))
    if subtitle:
        sp = doc.add_paragraph()
        _add_runs(sp, subtitle, size=12, color=TEXT_MUTED)

    for b in blocks:
        try:
            t = b["type"]
            if t == "heading":
                doc.add_heading("", level={1: 1, 2: 1, 3: 2, 4: 3}[min(4, max(1, b["level"]))])
                _add_runs(doc.paragraphs[-1], b["text"])
            elif t == "para":
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                _add_runs(p, b["text"])
            elif t in ("bullets", "numbered"):
                for level, text in b["items"]:
                    style = ("List Number" if t == "numbered" else "List Bullet") + (f" {level + 1}" if level else "")
                    try:
                        p = doc.add_paragraph(style=style)
                    except KeyError:
                        p = doc.add_paragraph(style="List Number" if t == "numbered" else "List Bullet")
                    _add_runs(p, text)
            elif t == "table":
                ncol = max(1, len(b["header"]))
                table = doc.add_table(rows=1, cols=ncol)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, htxt in enumerate(b["header"]):
                    cell = table.rows[0].cells[i]
                    cell.text = ""
                    _shade(cell, BRAND_DARK)
                    _add_runs(cell.paragraphs[0], htxt, size=10, color="#FFFFFF", bold=True)
                for r_i, row in enumerate(b["rows"], start=1):
                    cells = table.add_row().cells
                    for c_i, ctext in enumerate(row[:ncol]):
                        cells[c_i].text = ""
                        col = _status_color(ctext)
                        if col:
                            _shade(cells[c_i], col[0])
                            _add_runs(cells[c_i].paragraphs[0], ctext, size=10, color=col[1], bold=True)
                        else:
                            if r_i % 2 == 0:
                                _shade(cells[c_i], ZEBRA)
                            _add_runs(cells[c_i].paragraphs[0], ctext, size=10)
                doc.add_paragraph()
            elif t == "quote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                _add_runs(p, b["text"], color="#4B5563")
                for run in p.runs:
                    run.italic = True
            elif t == "code":
                p = doc.add_paragraph()
                run = p.add_run(_clean(b["text"])[:4000])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif t == "rule":
                doc.add_paragraph("_" * 60).alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif t == "pagebreak":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif t == "image" and b["name"] in images:
                doc.add_picture(io.BytesIO(images[b["name"]]), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b["alt"]:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr = cap.add_run(_clean(b["alt"]))
                    cr.italic = True
                    cr.font.size = Pt(9)
        except Exception as exc:
            HEALTH.flag("docx-build", f"block {b.get('type')}: {exc.__class__.__name__}: {exc}")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# =========================================================================================
# PPTX
# =========================================================================================
def compile_pptx_bytes(markdown_text: str, title: Optional[str] = None, subtitle: str = "",
                       images: Optional[Dict[str, bytes]] = None) -> bytes:
    """
    Markdown -> PowerPoint.  '# Title' = title slide; every '## Heading' = a slide; bullets / paragraphs / tables /
    images become slide content; a paragraph starting with 'Notes:' becomes speaker notes.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    images = images or {}
    blocks = parse_markdown(markdown_text)
    deck_title, blocks = extract_title(blocks, title or "Effiong AI Presentation")
    deck_title = _clean(title or deck_title)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    orange = RGBColor.from_string(BRAND_ORANGE.lstrip("#"))
    white = RGBColor(0xF0, 0xF6, 0xFC)
    muted = RGBColor(0x8B, 0x94, 0x9E)
    panel = RGBColor.from_string(BRAND_PANEL.lstrip("#"))

    def new_slide() -> Any:
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(BRAND_DARK.lstrip("#"))
        return s

    def textbox(slide: Any, x: float, y: float, w: float, h: float) -> Any:
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.word_wrap = True
        return tb.text_frame

    def chrome(slide: Any, number: int) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.32), prs.slide_width, Inches(0.18))
        bar.fill.solid()
        bar.fill.fore_color.rgb = orange
        bar.line.fill.background()
        tf = textbox(slide, 0.6, 6.9, 8, 0.35)
        tf.text = "EFFIONG AI"
        tf.paragraphs[0].runs[0].font.size = Pt(10)
        tf.paragraphs[0].runs[0].font.color.rgb = muted
        tf2 = textbox(slide, 11.7, 6.9, 1.2, 0.35)
        tf2.text = str(number)
        tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
        tf2.paragraphs[0].runs[0].font.size = Pt(10)
        tf2.paragraphs[0].runs[0].font.color.rgb = muted

    # title slide
    s = new_slide()
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.3), Inches(0.12), Inches(2.0))
    strip.fill.solid()
    strip.fill.fore_color.rgb = orange
    strip.line.fill.background()
    tf = textbox(s, 1.0, 2.2, 11.3, 1.6)
    tf.text = deck_title
    tf.paragraphs[0].runs[0].font.size = Pt(44)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = white
    sub = subtitle
    if not sub and blocks and blocks[0]["type"] == "para":
        sub = blocks[0]["text"]
        blocks = blocks[1:]
    tf = textbox(s, 1.0, 3.9, 11.3, 1.2)
    tf.text = _clean(plain(sub)) if sub else datetime.now().strftime("%d %B %Y")
    tf.paragraphs[0].runs[0].font.size = Pt(20)
    tf.paragraphs[0].runs[0].font.color.rgb = orange
    tf = textbox(s, 1.0, 6.3, 6, 0.4)
    tf.text = "Effiong AI - Sovereign Wisdom Engine"
    tf.paragraphs[0].runs[0].font.size = Pt(12)
    tf.paragraphs[0].runs[0].font.color.rgb = muted

    # split into slides by heading
    slides: List[Dict[str, Any]] = []
    cur: Optional[Dict[str, Any]] = None
    for b in blocks:
        if b["type"] == "heading" and b["level"] <= 2:
            cur = {"title": plain(b["text"]), "blocks": []}
            slides.append(cur)
        else:
            if cur is None:
                cur = {"title": "Overview", "blocks": []}
                slides.append(cur)
            cur["blocks"].append(b)

    number = 1
    for sd in slides:
        notes = []
        lines: List[Tuple[int, str, bool]] = []      # (level, text, is_bullet)
        tables, imgs = [], []
        for b in sd["blocks"]:
            if b["type"] in ("bullets", "numbered"):
                for k, (lv, tx) in enumerate(b["items"], 1):
                    lines.append((lv, (f"{k}. " if b["type"] == "numbered" else "") + tx, b["type"] == "bullets"))
            elif b["type"] == "para":
                if b["text"].lower().startswith("notes:"):
                    notes.append(b["text"][6:].strip())
                else:
                    lines.append((0, b["text"], False))
            elif b["type"] == "heading":
                lines.append((0, "**" + b["text"] + "**", False))
            elif b["type"] == "table":
                tables.append(b)
            elif b["type"] == "image" and b["name"] in images:
                imgs.append(b)
            elif b["type"] == "quote":
                lines.append((0, "*" + b["text"] + "*", False))
        # paginate long bullet lists
        per_slide = 7
        chunks = [lines[i:i + per_slide] for i in range(0, max(1, len(lines)), per_slide)] or [[]]
        for ci, chunk in enumerate(chunks):
            number += 1
            slide = new_slide()
            chrome(slide, number)
            tf = textbox(slide, 0.6, 0.45, 12.1, 1.0)
            tf.text = _clean(sd["title"]) + (" (cont.)" if ci else "")
            tf.paragraphs[0].runs[0].font.size = Pt(32)
            tf.paragraphs[0].runs[0].font.bold = True
            tf.paragraphs[0].runs[0].font.color.rgb = white
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.4), Inches(1.6), Inches(0.07))
            line.fill.solid()
            line.fill.fore_color.rgb = orange
            line.line.fill.background()

            body_w = 12.1 if not (imgs and ci == 0) else 6.9
            if chunk:
                total_chars = sum(len(t) for _, t, _ in chunk)
                size = 26 if total_chars < 380 else 20 if total_chars < 700 else 16
                tf = textbox(slide, 0.6, 1.75, body_w, 4.9)
                first = True
                for lv, tx, is_b in chunk:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    prefix = ("• " if lv == 0 else "– ") if is_b else ""
                    p.level = 0
                    p.space_after = Pt(9)
                    runs = parse_inline(_clean(tx))
                    for k, (rt, fl) in enumerate(runs):
                        r = p.add_run()
                        r.text = (prefix if k == 0 else "") + rt
                        r.font.size = Pt(size - (2 if lv else 0))
                        r.font.bold = True if fl.get("b") else None
                        r.font.italic = True if fl.get("i") else None
                        r.font.color.rgb = white if lv == 0 else RGBColor(0xC9, 0xD1, 0xD9)
                    if lv:
                        p.level = min(lv, 4)
            if ci == 0 and tables:
                tb = tables[0]
                nrow, ncol = len(tb["rows"]) + 1, max(1, len(tb["header"]))
                shape = slide.shapes.add_table(min(nrow, 9), ncol, Inches(0.6), Inches(3.4 if chunk else 1.8), Inches(12.1), Inches(0.4 * min(nrow, 9)))
                table = shape.table
                for c_i, h in enumerate(tb["header"]):
                    cell = table.cell(0, c_i)
                    cell.text = _clean(plain(h))
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = orange
                    cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
                    cell.text_frame.paragraphs[0].runs[0].font.bold = True
                    cell.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x0B, 0x0F, 0x17)
                for r_i, row in enumerate(tb["rows"][:8], start=1):
                    for c_i, val in enumerate(row[:ncol]):
                        cell = table.cell(r_i, c_i)
                        cell.text = _clean(plain(val))
                        cell.fill.solid()
                        col = _status_color(val)
                        cell.fill.fore_color.rgb = RGBColor.from_string(col[0].lstrip("#")) if col else (panel if r_i % 2 else RGBColor(0x16, 0x1B, 0x24))
                        if cell.text_frame.paragraphs[0].runs:
                            run = cell.text_frame.paragraphs[0].runs[0]
                            run.font.size = Pt(12)
                            run.font.color.rgb = RGBColor.from_string(col[1].lstrip("#")) if col else white
            if ci == 0 and imgs:
                try:
                    pic = slide.shapes.add_picture(io.BytesIO(images[imgs[0]["name"]]), Inches(7.8), Inches(1.75), width=Inches(4.9))
                    if pic.height > Inches(4.9):
                        ratio = Inches(4.9) / pic.height
                        pic.height, pic.width = int(pic.height * ratio), int(pic.width * ratio)
                except Exception as exc:
                    HEALTH.flag("pptx-build", f"image: {exc}")
            if notes and ci == 0:
                slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# =========================================================================================
# Charts
# =========================================================================================
PALETTE = ["#D27D2D", "#3498DB", "#2ECC71", "#9B59B6", "#E74C3C", "#F1C40F", "#1ABC9C", "#E67E22"]


def render_chart(spec: Dict[str, Any], theme: str = "dark") -> bytes:
    """
    spec = {"type": "bar|barh|line|area|pie|scatter", "title": str, "x_label": str, "y_label": str,
            "labels": [...], "series": [{"name": str, "values": [...]}, ...]}
    """
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    dark = theme == "dark"
    bg, fg, grid = ("#0B0F17", "#F0F6FC", "#212631") if dark else ("#FFFFFF", "#1F2937", "#E5E7EB")
    ctype = str(spec.get("type", "bar")).lower()
    labels = [str(x) for x in (spec.get("labels") or [])]
    series = spec.get("series") or []
    fig, ax = plt.subplots(figsize=(9, 5.2), dpi=140)
    fig.patch.set_facecolor(bg)
    ax.set_facecolor(bg)
    for spine in ax.spines.values():
        spine.set_color(grid)
    ax.tick_params(colors=fg)
    ax.title.set_color(fg)
    ax.xaxis.label.set_color(fg)
    ax.yaxis.label.set_color(fg)
    if ctype == "pie" and series:
        vals = [float(v) for v in series[0].get("values", [])]
        ax.pie(vals, labels=labels[:len(vals)], autopct="%1.0f%%", colors=PALETTE, textprops={"color": fg},
               wedgeprops={"edgecolor": bg, "linewidth": 2})
        ax.axis("equal")
    elif ctype == "scatter":
        for i, s in enumerate(series):
            xs = s.get("x") or list(range(len(s.get("values", []))))
            ax.scatter(xs, s.get("values", []), color=PALETTE[i % len(PALETTE)], label=s.get("name"), s=48)
    else:
        n = max((len(s.get("values", [])) for s in series), default=0)
        xs = list(range(n))
        if not labels:
            labels = [str(i + 1) for i in xs]
        width = 0.8 / max(1, len(series))
        for i, s in enumerate(series):
            vals = [float(v) for v in s.get("values", [])]
            color = PALETTE[i % len(PALETTE)]
            if ctype == "line":
                ax.plot(xs[:len(vals)], vals, marker="o", color=color, label=s.get("name"), linewidth=2.4)
            elif ctype == "area":
                ax.fill_between(xs[:len(vals)], vals, alpha=0.35, color=color)
                ax.plot(xs[:len(vals)], vals, color=color, label=s.get("name"), linewidth=2)
            elif ctype == "barh":
                ax.barh([x + i * width for x in xs[:len(vals)]], vals, height=width, color=color, label=s.get("name"))
            else:
                ax.bar([x + i * width for x in xs[:len(vals)]], vals, width=width, color=color, label=s.get("name"))
        shift = width * (len(series) - 1) / 2 if ctype in ("bar", "barh") else 0
        if ctype == "barh":
            ax.set_yticks([x + shift for x in xs])
            ax.set_yticklabels(labels[:n])
        else:
            ax.set_xticks([x + shift for x in xs])
            ax.set_xticklabels(labels[:n], rotation=30 if n > 6 else 0, ha="right" if n > 6 else "center")
        ax.grid(color=grid, linestyle="--", linewidth=0.6, alpha=0.8)
        ax.set_axisbelow(True)
    ax.set_title(str(spec.get("title", "")), fontsize=14, fontweight="bold", pad=12)
    if spec.get("x_label") and ctype != "pie":
        ax.set_xlabel(str(spec["x_label"]))
    if spec.get("y_label") and ctype != "pie":
        ax.set_ylabel(str(spec["y_label"]))
    if len([s for s in series if s.get("name")]) > 1 and ctype != "pie":
        leg = ax.legend(facecolor=bg, edgecolor=grid)
        for t in leg.get_texts():
            t.set_color(fg)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=bg)
    plt.close(fig)
    return buf.getvalue()


# =========================================================================================
# SVG (vector graphics) sanitiser
# =========================================================================================
_SVG_NS = "http://www.w3.org/2000/svg"
_BLOCKED_TAGS = {"script", "foreignobject", "iframe", "object", "embed", "audio", "video", "canvas", "link", "meta", "use_external"}


def sanitize_svg(svg: str) -> Optional[str]:
    """Return a safe SVG string (no scripts, no event handlers, no external references) or None if unusable."""
    if not svg:
        return None
    m = re.search(r"<svg[\s\S]*</svg>", svg, re.IGNORECASE)
    if not m:
        return None
    src = re.sub(r"<!DOCTYPE[^>]*>|<!ENTITY[^>]*>|<\?xml[^>]*\?>", "", m.group(0))
    ET.register_namespace("", _SVG_NS)
    ET.register_namespace("xlink", "http://www.w3.org/1999/xlink")
    try:
        root = ET.fromstring(src)
    except ET.ParseError:
        return None

    def clean(el: ET.Element) -> None:
        for child in list(el):
            tag = child.tag.split("}")[-1].lower()
            if tag in _BLOCKED_TAGS:
                el.remove(child)
                continue
            clean(child)
        for attr in list(el.attrib):
            name = attr.split("}")[-1].lower()
            val = str(el.attrib[attr]).strip().lower()
            if name.startswith("on"):
                del el.attrib[attr]
            elif name == "href" and not (val.startswith("#") or val.startswith("data:image/")):
                del el.attrib[attr]
            elif "javascript:" in val:
                del el.attrib[attr]
            elif name == "style" and ("url(http" in val or "@import" in val or "expression(" in val):
                del el.attrib[attr]
        if el.tag.split("}")[-1].lower() == "style" and el.text:
            el.text = re.sub(r"@import[^;]*;|url\(\s*['\"]?https?:[^)]*\)", "", el.text)

    clean(root)
    if not root.get("xmlns") and root.tag == "svg":
        root.set("xmlns", _SVG_NS)
    if not root.get("viewBox") and root.get("width") and root.get("height"):
        try:
            root.set("viewBox", f"0 0 {float(re.sub('[^0-9.]', '', root.get('width')))} {float(re.sub('[^0-9.]', '', root.get('height')))}")
        except ValueError:
            pass
    return ET.tostring(root, encoding="unicode")
